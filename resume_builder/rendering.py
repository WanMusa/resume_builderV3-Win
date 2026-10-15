"""Single-column DOCX rendering. Fixed facts are taken only from the base."""

import hashlib
from datetime import datetime
import os
from pathlib import Path
import re
import uuid

from docx import Document
from docx.shared import Inches, Pt, RGBColor

from .models import BaseResume
from .providers import validate_content


class RenderError(ValueError):
    """Safe, actionable publication failure."""


def _document(base, mock):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.name = "Aptos"
        doc.styles[name].font.color.rgb = RGBColor.from_string("174A73")
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    if mock:
        doc.add_paragraph("NOT FOR SUBMISSION — MOCK DOCUMENT", "Heading 1")
        section.header.paragraphs[0].text = "NOT FOR SUBMISSION — MOCK DOCUMENT"
    doc.add_paragraph(base.identity.name, "Title")
    contact = [base.identity.email, base.identity.phone, base.identity.location, *base.identity.links]
    if any(contact):
        doc.add_paragraph(" | ".join(x for x in contact if x))
    return doc


def _resume(base, content, mock):
    doc = _document(base, mock)
    doc.add_heading("Career Profile", level=1)
    doc.add_paragraph(content["summary"]["text"])
    if content["expertise"]:
        doc.add_heading("Expertise", level=1)
        doc.add_paragraph(" • ".join(content["expertise"]))
    if base.roles:
        doc.add_heading("Experience", level=1)
    for role in base.roles:
        doc.add_heading(role.title, level=2)
        doc.add_paragraph(" | ".join(x for x in (role.employer, role.dates) if x))
        if role.background:
            doc.add_paragraph(role.background)
        for achievement in content["achievements"]:
            if achievement["role_id"] == role.id:
                doc.add_paragraph(achievement["text"], "List Bullet")
    if base.education:
        doc.add_heading("Education", level=1)
        for entry in base.education:
            doc.add_paragraph(" | ".join(x for x in (entry.qualification, entry.institution, entry.dates) if x))
            if entry.details:
                doc.add_paragraph(entry.details)
    if base.certifications:
        doc.add_heading("Certifications", level=1)
        for entry in base.certifications:
            doc.add_paragraph(entry, "List Bullet")
    if base.references:
        doc.add_heading("References", level=1)
        doc.add_paragraph(base.references)
    for section in base.additional_sections:
        doc.add_heading(section.heading, level=1)
        for entry in section.entries:
            doc.add_paragraph(entry)
    return doc


def _letter(base, content, mock, job):
    doc = _document(base, mock)
    if job.get("title"):
        doc.add_heading("Application: " + job["title"], level=1)
    if job.get("company"):
        doc.add_paragraph(job["company"])
    captured_at = job.get("captured_at")
    if captured_at:
        try:
            captured_date = datetime.fromisoformat(captured_at.replace("Z", "+00:00")).date()
            doc.add_paragraph(captured_date.isoformat())
        except (ValueError, TypeError, AttributeError):
            pass
    doc.add_paragraph("Dear Hiring Manager,")
    for paragraph in content["paragraphs"]:
        doc.add_paragraph(paragraph["text"])
    doc.add_paragraph("Sincerely,")
    doc.add_paragraph(base.identity.name)
    return doc


def render_documents(job_id: str, revision: int, job: dict, base: dict, resume: dict,
                     letter: dict, output_folder: str, mock: bool) -> list[str]:
    """Publish without replacing existing files, including user-edited documents."""
    if type(revision) is not int or revision < 1 or revision > 1_000_000:
        raise RenderError("Revision must be an integer between 1 and 1,000,000.")
    if not isinstance(job_id, str) or not job_id or len(job_id) > 500:
        raise RenderError("Job ID is missing or too long.")
    base_model = BaseResume.model_validate(base)
    resume = validate_content("resume", resume, base)
    letter = validate_content("letter", letter, base)
    slug = re.sub(r"[^a-zA-Z0-9_-]", "-", job_id).strip("-")[:40] or "job"
    digest = hashlib.sha256(job_id.encode()).hexdigest()[:12]
    if not isinstance(output_folder, str) or not output_folder.strip():
        raise RenderError("Choose an output folder.")
    root = Path(output_folder).expanduser().absolute()
    labels = [
        re.sub(r"[^a-zA-Z0-9_-]", "-", str(job.get(field) or "")[:120]).strip("-")
        for field in ("title", "company")
    ]
    labels = [label for label in labels if label]
    if labels:
        # Keep room for the stable job discriminator and revision on Windows.
        suffix = f"-{slug[:12]}-{digest}-r{revision}"
        available = min(48, 200 - len(str(root)) - 1 - len(suffix))
        if available < 8:
            raise RenderError("Output path is too long for portable Word files; choose a shorter folder.")
        per_label = (available - len(labels) + 1) // len(labels)
        directory = "-".join(label[:per_label].rstrip("-") for label in labels) + suffix
    else:
        directory = f"{slug}-{digest}-r{revision}"
    target = root / directory
    if len(str(target)) > 200:
        raise RenderError("Output path is too long for portable Word files; choose a shorter folder.")
    pending, published = [], []
    try:
        root.mkdir(parents=True, exist_ok=True)
        if target.is_symlink():
            raise RenderError("Output revision directory cannot be a symbolic link.")
        target.mkdir(exist_ok=True)
        destinations = [target / "resume.docx", target / "cover-letter.docx"]
        if any(p.exists() or p.is_symlink() for p in destinations):
            raise RenderError("Outputs already exist and will not be overwritten. Keep edits and create a new revision.")
        docs = [_resume(base_model, resume, mock), _letter(base_model, letter, mock, job)]
        for doc in docs:
            path = target / (uuid.uuid4().hex + ".pending")
            pending.append(path)
            with path.open("xb") as stream:
                doc.save(stream)
                stream.flush()
                os.fsync(stream.fileno())
        for source, destination in zip(pending, destinations):
            # A hard link atomically creates a new filename, never replacing it.
            os.link(source, destination)
            published.append(destination)
        return [str(p) for p in destinations]
    except FileExistsError:
        raise RenderError("Outputs already exist; create a new revision rather than overwrite edited documents.") from None
    except PermissionError:
        raise RenderError("Cannot publish DOCX files. Close Word if files are locked and check output folder permissions.") from None
    except OSError:
        raise RenderError("Cannot publish DOCX files. Check disk space, folder permissions and Word locks; use a local folder supporting atomic file creation.") from None
    finally:
        for path in pending:
            path.unlink(missing_ok=True)
        # Roll back only files this invocation published if the pair was incomplete.
        if len(published) != 2:
            for path in published:
                path.unlink(missing_ok=True)
