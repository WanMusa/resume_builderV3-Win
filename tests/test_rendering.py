from pathlib import Path
import shutil
import uuid

from docx import Document
import pytest

from resume_builder import rendering
from test_providers import base, mock_result


@pytest.fixture
def output_folder():
    # Keep all test file operations inside the project, not system temp folders.
    path = Path.cwd() / ".test-output" / uuid.uuid4().hex
    path.mkdir(parents=True)
    yield path
    shutil.rmtree(path)
    try:
        path.parent.rmdir()
    except OSError:
        pass


def render(base, folder, **kwargs):
    return rendering.render_documents(
        kwargs.get("job_id", "../unsafe:job"), kwargs.get("revision", 1), kwargs.get("job", {}), base,
        mock_result(base)["content"], mock_result(base, "letter")["content"],
        str(folder), kwargs.get("mock", True),
    )


def text(path):
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs)


def test_fixed_facts_and_single_column_mock_documents(base, output_folder):
    paths = render(base, output_folder)
    assert len(paths) == 2
    resume = text(paths[0])
    for fixed in ("Fixed Person", "Fixed Employer", "Fixed Title", "2020–2024",
                  "Fixed Degree", "Fixed School", "Fixed Certification", "Fixed volunteering",
                  "Available on request"):
        assert fixed in resume
    for path in paths:
        document = Document(path)
        assert "NOT FOR SUBMISSION" in text(path)
        assert "NOT FOR SUBMISSION" in document.sections[0].header.paragraphs[0].text
        assert not document.tables
        assert document.styles["Normal"].font.name == "Aptos"
        assert document.core_properties.author == ""
        assert Path(path).resolve().is_relative_to(output_folder.resolve())
    assert not list(output_folder.rglob("*.pending"))


def test_does_not_overwrite_edited_outputs(base, output_folder):
    paths = render(base, output_folder)
    Path(paths[0]).write_bytes(b"user edits")
    with pytest.raises(rendering.RenderError, match="not be overwritten"):
        render(base, output_folder)
    assert Path(paths[0]).read_bytes() == b"user edits"
    other = render(base, output_folder, revision=2)
    assert Path(other[0]).parent != Path(paths[0]).parent


def test_safe_distinct_job_directories(base, output_folder):
    first = render(base, output_folder, job_id="a/b")
    second = render(base, output_folder, job_id="a?b")
    assert Path(first[0]).parent != Path(second[0]).parent


def test_word_lock_error_is_actionable(base, output_folder, monkeypatch):
    monkeypatch.setattr(rendering.os, "link", lambda *a: (_ for _ in ()).throw(PermissionError("private")))
    with pytest.raises(rendering.RenderError, match="Close Word") as error:
        render(base, output_folder)
    assert "private" not in str(error.value)
    assert not list(output_folder.rglob("*.pending"))
    assert not list(output_folder.rglob("*.docx"))


def test_incomplete_pair_is_rolled_back(base, output_folder, monkeypatch):
    original = rendering.os.link
    calls = []

    def publish(source, destination):
        calls.append(destination)
        if len(calls) == 2:
            raise PermissionError()
        original(source, destination)

    monkeypatch.setattr(rendering.os, "link", publish)
    with pytest.raises(rendering.RenderError):
        render(base, output_folder)
    assert not list(output_folder.rglob("*.docx"))
    assert not list(output_folder.rglob("*.pending"))


def test_long_paths_and_revision_validation(base, output_folder):
    with pytest.raises(rendering.RenderError, match="too long"):
        render(base, str(output_folder / ("a" * 180)))
    with pytest.raises(rendering.RenderError, match="Revision"):
        render(base, output_folder, revision=0)


def test_job_labels_in_directory_and_letter_come_from_source(base, output_folder):
    job = {"title": "Quality Engineer", "company": "Fictional Company",
           "captured_at": "2026-10-15T08:00:00+00:00"}
    paths = render(base, output_folder, job_id="uuid-test-job", revision=3, job=job)
    directory = Path(paths[0]).parent.name
    assert "Quality-Engineer" in directory
    assert "Fictional-Company" in directory
    assert directory.endswith("-r3")
    letter = text(paths[1])
    assert "Application: Quality Engineer" in letter
    assert "Fictional Company" in letter
    assert "2026-10-15" in letter


def test_job_directory_labels_are_sanitized_and_bounded(base, output_folder):
    job = {"title": "../" + "Role" * 200, "company": "Bad:/Name" * 100}
    paths = render(base, output_folder, job=job)
    assert len(str(Path(paths[0]).parent)) <= 200
    assert Path(paths[0]).resolve().is_relative_to(output_folder.resolve())
    assert ":" not in Path(paths[0]).parent.name
